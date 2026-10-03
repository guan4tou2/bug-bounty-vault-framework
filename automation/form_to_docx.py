#!/usr/bin/env python3
"""FORM markdown -> vendor-form .docx (reference implementation).

Many programs want each finding on the vendor's own Word intake form (a two-column
table). Producing those by hand means the .md source and the .docx drift. This script
reads the FORM markdown (the linted source of truth) and fills a copy of the vendor's
Word template, so the two cannot drift.

Formatting is cloned from an existing accepted .docx (font, borders, section shading) by
copying its <w:tr> elements — far more faithful than rebuilding them through python-docx.

**You provide the template** (`--template your-vendor-form.docx`); there is no default,
because the template IS the vendor-specific part. The markdown conventions this assumes:
  - `# Title` document title
  - `## Section` headings; a section with `### Sub` subheadings renders as prose rows,
    otherwise as a `| label | value |` key/value table
  - one column embeds images: the field whose label contains the image keyword
    (default "screenshot"; override with --image-field or BB_IMAGE_FIELD)
Adapt those for your own form if it differs.

Usage:
    python3 automation/form_to_docx.py <FORM.md> [...] --template <ref.docx> --out <dir>
    python3 automation/form_to_docx.py <FORM.md> --template f.docx --out d --prefix "Report - "
"""
from __future__ import annotations

import argparse
import copy
import os
import re
import sys
from pathlib import Path

try:
    import docx
except ImportError:  # pragma: no cover
    sys.exit("python-docx not installed — `uv pip install python-docx`")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Label column 25% / value 75% of a 9072 dxa table width.
LABEL_W, VALUE_W = 2268, 6804

# Header-row labels to skip in a key/value table (first cell). Override: BB_FORM_HEADERS.
HEADER_LABELS = set(
    (os.environ.get("BB_FORM_HEADERS") or "field|label|key|item").split("|"))
# Which field embeds images (substring match on the label). Override: BB_IMAGE_FIELD.
IMAGE_FIELD_KW = (os.environ.get("BB_IMAGE_FIELD") or "screenshot").lower()


# ---------------------------------------------------------------- markdown side

def strip_meta(src: str) -> str:
    src = re.sub(r"\A---\n.*?\n---\n", "", src, flags=re.S)
    return re.sub(r"<!--.*?-->", "", src, flags=re.S)


def split_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_kv_table(block: str) -> list[tuple[str, str]]:
    """The key/value tables of the non-prose sections.

    Prose that sits OUTSIDE the table (e.g. a blockquote appended after it) is kept as a
    trailing row rather than dropped — silently losing a paragraph of a report is the
    worst possible failure mode for this script.
    """
    out: list[tuple[str, str]] = []
    trailing: list[str] = []
    for line in block.splitlines():
        line = line.rstrip()
        s = line.strip()
        if s.startswith("|"):
            cells = split_table_row(s)
            if len(cells) < 2:
                continue
            if re.fullmatch(r"[-: ]+", cells[0]) or cells[0].lower() in HEADER_LABELS:
                continue
            out.append((cells[0], cells[1]))
        elif s:
            trailing.append(line)
    if trailing:
        text = "\n".join(trailing).strip()
        label = "Note"
        m = re.match(r"^>?\s*\*\*(.{2,20}?)(\([^)]*\))?\*\*", text)
        if m:
            label = m.group(1)
        out.append((label, text))
    return out


def parse_form(md: str) -> tuple[list[tuple[str, list]], str]:
    """-> ([(section_title, [(label, value_markdown)])], document title)"""
    md = strip_meta(md)
    doc_title = ""
    m = re.search(r"^# +(.+)$", md, re.M)
    if m:
        doc_title = m.group(1).strip()

    sections: list[tuple[str, list]] = []
    parts = re.split(r"^## +(.+)$", md, flags=re.M)[1:]
    for title, body in zip(parts[0::2], parts[1::2]):
        title = title.strip()
        # Prose-mode section: has ### subheadings (language-agnostic, replaces a
        # vendor-specific section-name check).
        if re.search(r"^### +.+$", body, flags=re.M):
            rows = []
            subs = re.split(r"^### +(.+)$", body, flags=re.M)
            lead = subs[0].strip()
            if lead:
                rows.append(("Details", lead))
            for sub_t, sub_b in zip(subs[1::2], subs[2::2]):
                rows.append((sub_t.strip(), sub_b.strip()))
            sections.append((title, rows))
        else:
            sections.append((title, parse_kv_table(body)))
    return sections, doc_title


def to_lines(value: str) -> list[tuple[str, bool]]:
    """Flatten a cell's markdown into (text, is_monospace) display lines.

    Nested tables become `a | b` lines: Word tables inside table cells render badly and
    the reader only needs the pairing.
    """
    value = value.replace("<br>", "\n")
    lines: list[tuple[str, bool]] = []
    in_code = False
    for raw in value.split("\n"):
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            lines.append((line, True))
            continue
        s = line.strip()
        if not s:
            if lines and lines[-1][0] != "":
                lines.append(("", False))
            continue
        if s.startswith("|"):
            cells = split_table_row(s)
            if all(re.fullmatch(r"[-: ]*", c) for c in cells):
                continue
            lines.append((" | ".join(c for c in cells if c != ""), False))
            continue
        s = re.sub(r"^- \[[ x]\] *", "> ", s)
        s = re.sub(r"^[-*] +", "- ", s)
        s = re.sub(r"^> *", "", s)
        lines.append((s, False))
    while lines and lines[-1][0] == "":
        lines.pop()
    return lines


def inline_segments(text: str) -> list[tuple[str, bool, bool]]:
    """-> [(text, bold, mono)] from **bold** and `code`."""
    segs: list[tuple[str, bool, bool]] = []
    for chunk in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            for sub, _, sub_mono in inline_segments(chunk[2:-2]):
                segs.append((sub, True, sub_mono))
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            segs.append((chunk[1:-1], False, True))
        else:
            segs.append((chunk, False, False))
    return segs or [(text, False, False)]


# ------------------------------------------------------------------ docx side

class FormBuilder:
    def __init__(self, template: Path):
        self.doc = docx.Document(str(template))
        self.table = self.doc.tables[0]
        rows = self.table.rows
        self.hdr_tr = copy.deepcopy(rows[0]._tr)
        self.row_tr = copy.deepcopy(rows[1]._tr)
        self.hdr_p = copy.deepcopy(rows[0].cells[0].paragraphs[0]._p)
        self.lbl_p = copy.deepcopy(rows[1].cells[0].paragraphs[0]._p)
        self.val_p = copy.deepcopy(rows[1].cells[1].paragraphs[0]._p)
        for tr in list(self.table._tbl.findall(W + "tr")):
            self.table._tbl.remove(tr)

    @staticmethod
    def _runs_of(p):
        return p.findall(W + "r")

    def _fill(self, tc, p_proto, lines: list[tuple[str, bool]], force_bold=False):
        for p in list(tc.findall(W + "p")):
            tc.remove(p)
        proto_run = None
        for r in self._runs_of(p_proto):
            proto_run = r
            break
        for text, mono in lines or [("", False)]:
            p = copy.deepcopy(p_proto)
            for r in self._runs_of(p):
                p.remove(r)
            for seg, bold, seg_mono in inline_segments(text):
                r = copy.deepcopy(proto_run)
                rPr = r.find(W + "rPr")
                if rPr is not None:
                    for tag in ("b", "bCs"):
                        for el in rPr.findall(W + tag):
                            rPr.remove(el)
                    if bold or force_bold:
                        b = rPr.makeelement(W + "b", {})
                        rPr.insert(0, b)
                    if mono or seg_mono:
                        fonts = rPr.find(W + "rFonts")
                        if fonts is not None:
                            for a in ("ascii", "hAnsi"):
                                fonts.set(W + a, "Consolas")
                t = r.find(W + "t")
                if t is None:
                    t = r.makeelement(W + "t", {})
                    r.append(t)
                t.text = seg
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                p.append(r)
            tc.append(p)

    @staticmethod
    def _set_width(tc, dxa: int):
        tcPr = tc.find(W + "tcPr")
        if tcPr is None:
            return
        tcW = tcPr.find(W + "tcW")
        if tcW is not None:
            tcW.set(W + "w", str(dxa))

    def add_section(self, title: str):
        tr = copy.deepcopy(self.hdr_tr)
        tc = tr.findall(W + "tc")[0]
        self._fill(tc, self.hdr_p, [(title, False)], force_bold=True)
        self.table._tbl.append(tr)

    def add_row(self, label: str, value: str, images: "list[Path] | None" = None,
                captions: "dict[str, str] | None" = None):
        tr = copy.deepcopy(self.row_tr)
        tcs = tr.findall(W + "tc")
        self._set_width(tcs[0], LABEL_W)
        self._set_width(tcs[1], VALUE_W)
        self._fill(tcs[0], self.lbl_p, [(label, False)], force_bold=True)
        # The image field's text body is a "filename — caption" list; since the images are
        # embedded, that filename list is noise to the reader (and looks like it points at
        # a non-existent attachment). Captions move under each image instead.
        body = "" if images else value
        self._fill(tcs[1], self.val_p, to_lines(body))
        self.table._tbl.append(tr)
        if images:
            self._embed(tcs[1], images, captions or {})

    def _embed(self, tc, images: "list[Path]", captions: "dict[str, str]"):
        """Embed screenshots directly in the field rather than attaching a folder.

        Embedding keeps the reviewer from having to map loose files back to fields. Width
        is the value-column width (~4.55 in) with a little margin so nothing is clipped.
        """
        from docx.shared import Inches
        from docx.table import _Cell
        cell = _Cell(tc, self.table)
        for img in images:
            if not img.exists():
                print(f"   ⚠ screenshot missing, skipped: {img.name}", file=sys.stderr)
                continue
            para = cell.add_paragraph()
            shape = para.add_run().add_picture(str(img), width=Inches(4.55))
            text = captions.get(img.name, "")
            # python-docx writes the filename into the image object name (<pic:cNvPr
            # name=…>). Invisible in the body, but visible in Word's alt-text pane and in
            # PDF accessibility metadata — so with the body no longer printing filenames,
            # leaving it here would be the same leak in another place.
            for attr, val in (("name", text or "PoC"), ("descr", text)):
                if val:
                    shape._inline.docPr.set(attr, val)
            try:
                pic = shape._inline.graphic.graphicData.pic
                pic.nvPicPr.cNvPr.set("name", text or "PoC")
                if text:
                    pic.nvPicPr.cNvPr.set("descr", text)
            except AttributeError:
                pass
            if text:
                cap = cell.add_paragraph()
                run = cap.add_run(text)
                run.font.size = __import__("docx").shared.Pt(8)

    def save(self, path: Path):
        self.doc.save(str(path))
        strip_unused_media(path)


def strip_unused_media(path: Path) -> None:
    """Drop images and the preview thumbnail inherited from the template.

    Cloning a real submitted .docx also clones word/media/ — so a generated report can
    ship a screenshot belonging to a *different* finding, plus a thumbnail of that other
    report's first page. That is a disclosure bug, not a size problem: it puts one
    customer's evidence inside another report. Anything not referenced by the rebuilt
    document.xml is removed here, along with the relationship and content-type entries.
    """
    import zipfile
    from xml.etree import ElementTree as ET

    PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"

    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}

    doc_xml = parts["word/document.xml"].decode("utf-8")
    used_rids = set(re.findall(r'r:(?:id|embed|link)="([^"]+)"', doc_xml))

    drop: set[str] = {n for n in parts if n.startswith("docProps/thumbnail")}

    rels_name = "word/_rels/document.xml.rels"
    rels = ET.fromstring(parts[rels_name])
    changed = False
    for rel in list(rels):
        target = rel.get("Target", "")
        if not target.startswith("media/"):
            continue
        if rel.get("Id") in used_rids:
            continue
        drop.add("word/" + target.lstrip("/"))
        rels.remove(rel)
        changed = True
    if changed:
        ET.register_namespace("", PKG.strip("{}"))
        parts[rels_name] = ET.tostring(rels, encoding="UTF-8", xml_declaration=True)

    if not drop:
        return

    root_rels = ET.fromstring(parts["_rels/.rels"])
    for rel in list(root_rels):
        if rel.get("Target", "").lstrip("/").startswith("docProps/thumbnail"):
            root_rels.remove(rel)
    ET.register_namespace("", PKG.strip("{}"))
    parts["_rels/.rels"] = ET.tostring(root_rels, encoding="UTF-8", xml_declaration=True)

    ctypes = ET.fromstring(parts["[Content_Types].xml"])
    for el in list(ctypes):
        name = el.get("PartName", "").lstrip("/")
        if name and name in drop:
            ctypes.remove(el)
    ET.register_namespace("", CT.strip("{}"))
    parts["[Content_Types].xml"] = ET.tostring(ctypes, encoding="UTF-8", xml_declaration=True)

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            if name not in drop:
                z.writestr(name, data)


def out_name(md_path: Path, prefix: str) -> str:
    stem = re.sub(r"^FORM - ", "", md_path.stem)
    return f"{prefix}{stem}.docx"


SHOT_LINE_RE = re.compile(
    r"^\s*[-*]\s*`?([A-Za-z0-9][A-Za-z0-9._-]*\.png)`?\s*(?:[—–-]\s*)?(.*)$")


def screenshot_captions(md: str) -> dict[str, str]:
    """filename -> human caption.

    The image field is written as ``- `shot-01.png` — caption``. The filename is for THIS
    script to locate the image, not for the reader: the image is embedded, so that .png is
    not attached — naming it points at a non-existent attachment. The filename stays in the
    md and out of the docx; only the caption (after the dash) travels with the image.
    """
    out: dict[str, str] = {}
    for line in md.splitlines():
        m = SHOT_LINE_RE.match(line)
        if m:
            out[m.group(1)] = re.sub(r"[`*]", "", m.group(2)).strip()
    return out


def find_screenshots(md_path: Path, captions: "dict[str, str] | None" = None) -> "list[Path]":
    """Resolve the filenames named in the FORM's image field to paths under Evidence/.

    Deliberately only trusts filenames WRITTEN in the FORM; it does not glob Evidence/.
    Globbing pulls in another finding's screenshots (an early version cloned the template
    and embedded another finding's image into every report — cross-case evidence leak, not
    a size issue). What the FORM names is what ships; a mismatch warns, never guesses.
    """
    md = md_path.read_text(encoding="utf-8")
    captions = captions if captions is not None else {}
    names = re.findall(r"([A-Za-z0-9][A-Za-z0-9._-]*\.png)", md)
    if not names:
        return []
    captions.update(screenshot_captions(md))
    # The target root is the dir that has both Evidence/ and Findings/ — Evidence/ alone
    # would stop at a stray Submissions/Forms/Evidence and then report "screenshot missing".
    tdir = md_path.resolve()
    for _ in range(6):
        tdir = tdir.parent
        if (tdir / "Evidence").is_dir() and (tdir / "Findings").is_dir():
            break
    ev = tdir / "Evidence"
    out, missing = [], []
    for n in dict.fromkeys(names):
        pth = ev / n
        (out if pth.exists() else missing).append(pth)
    for m in missing:
        print(f"   ⚠ {md_path.name} names {m.name} but it is not under Evidence/", file=sys.stderr)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("forms", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--template", type=Path, required=True,
                    help="your vendor's Word form (.docx) to clone formatting from")
    ap.add_argument("--prefix", default="Report - ",
                    help="output filename prefix (default: 'Report - ')")
    args = ap.parse_args()

    if not args.template.is_file():
        sys.exit(f"--template not found: {args.template}\n"
                 "Provide your vendor's Word form; there is no default (the template IS "
                 "the vendor-specific part).")

    args.out.mkdir(parents=True, exist_ok=True)
    for md_path in args.forms:
        sections, _ = parse_form(md_path.read_text(encoding="utf-8"))
        if not sections:
            print(f"!! {md_path.name}: no ## sections found", file=sys.stderr)
            continue
        caps: dict[str, str] = {}
        shots = find_screenshots(md_path, caps)
        b = FormBuilder(args.template)
        for title, rows in sections:
            b.add_section(title)
            for label, value in rows:
                imgs = shots if IMAGE_FIELD_KW in label.lower() else None
                b.add_row(label, value, images=imgs, captions=caps)
        dest = args.out / out_name(md_path, args.prefix)
        b.save(dest)
        print(f"OK {dest.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
